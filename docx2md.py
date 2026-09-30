#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""docx2md — 将 Word (.docx) 文件转换为 Markdown 的单文件独立脚本

用法:
    python docx2md path/to/file.docx

运行后自动在 docx 文件所在目录生成同名 .md 文件。

依赖:
    pip install python-docx

"""

import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from docx import Document
from docx.oxml.ns import qn
from docx.styles import BabelFish


# ══════════════════════════════════════════════════════════════════
# Block 数据结构
# ══════════════════════════════════════════════════════════════════

@dataclass
class Block:
    """内容块，按 y 坐标排列"""
    type: str  # "paragraph" / "table" / "image"
    y_top: float = 0.0
    text: str = ""
    headers: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    image_path: str = ""


# ══════════════════════════════════════════════════════════════════
# ConvertOptions（docx 转换相关子集）
# ══════════════════════════════════════════════════════════════════

@dataclass
class ConvertOptions:
    """转换行为开关"""
    remove_cjk_spaces: bool = True
    merge_broken_paragraphs: bool = True
    fix_headings: bool = True
    merge_blank_lines: bool = True
    add_front_matter: bool = False
    business_mode: bool = False
    fix_list_number_positions: bool = False
    infer_empty_headers: bool = False
    merge_cross_page_tables: bool = True
    filter_footer: bool = True
    extract_images: bool = False
    enable_ocr: bool = True
    ocr_lang: str = "ch"
    page_range: Optional[tuple] = None

    @classmethod
    def default(cls) -> "ConvertOptions":
        return cls()

    @classmethod
    def business(cls) -> "ConvertOptions":
        opts = cls()
        opts.business_mode = True
        opts.fix_list_number_positions = True
        opts.infer_empty_headers = True
        return opts


def _opts(options: Optional[ConvertOptions]) -> ConvertOptions:
    return options if options is not None else ConvertOptions.default()


# ══════════════════════════════════════════════════════════════════
# 后处理（postprocess）
# ══════════════════════════════════════════════════════════════════

def remove_cjk_spaces(text: str) -> str:
    """去除中文字符之间的空格（半角与全角）"""
    text = re.sub(r'(?<=[\u4e00-\u9fff]) +(?=[\u4e00-\u9fff])', '', text)
    text = re.sub(r'(?<=[\u4e00-\u9fff])\u3000+(?=[\u4e00-\u9fff])', '', text)
    return text


def merge_broken_paragraphs(blocks, options: Optional[ConvertOptions] = None):
    """过滤无用内容 + 合并断裂短行 + 跨页续行合并 + 去除中文间多余空格"""
    opts = _opts(options)
    result = []
    for b in blocks:
        if b.type == "paragraph":
            text = b.text.strip()
            if re.match(r'^\d+$', text):
                continue
            if re.match(r'^[IVXLCDM]{1,6}$', text):
                continue
            if re.match(r'^[\.\s]{3,}$', text):
                continue

            cleaned = remove_cjk_spaces(text) if opts.remove_cjk_spaces else text
            b.text = cleaned

            if (opts.merge_broken_paragraphs
                    and result
                    and result[-1].type == "paragraph"
                    and len(cleaned) <= 6
                    and re.search(r'\d+\.\s*$', result[-1].text.strip())):
                result[-1].text += cleaned
                continue

            if (opts.merge_broken_paragraphs
                    and result
                    and result[-1].type == "paragraph"
                    and _is_cross_page_continuation(result[-1].text, cleaned)):
                result[-1].text += cleaned
                continue

            result.append(b)
        else:
            result.append(b)
    return result


_SENTENCE_END_RE = re.compile(r'[。！？；;…」』\)\)]$')
_PARAGRAPH_START_RE = re.compile(
    r'^（\d+[)）]|^\d+[.、]\s|^第\d+章|^【|^序号|^说明[:：]|^注[:：]|^#{1,6}\s|^\d+\.\d+'
)
_HEADING_NUM_PREFIX_RE = re.compile(r'^\d+(?:\.\d+)+\s')


def _is_cross_page_continuation(prev_text: str, curr_text: str) -> bool:
    prev = prev_text.strip()
    curr = curr_text.strip()
    if len(prev) < 10 or len(curr) < 10:
        return False
    if _SENTENCE_END_RE.search(prev):
        return False
    if _PARAGRAPH_START_RE.match(curr):
        return False
    if re.match(r'^#{1,6}\s', prev):
        return False
    if re.match(r'^第\d+章', prev):
        return False
    if _HEADING_NUM_PREFIX_RE.match(prev):
        return False
    return True


def fix_list_number_positions(blocks, options: Optional[ConvertOptions] = None):
    """修复列表序号位置（特定模式规则）"""
    opts = _opts(options)
    if not opts.fix_list_number_positions:
        return blocks

    for b in blocks:
        if b.type != "paragraph":
            continue

        text = b.text

        text = re.sub(r'(?<=[\u4e00-\u9fff])[ 　]*\d+\.\s*(?=[\u4e00-\u9fff])', '', text)
        b.text = text

        m = re.match(r'^([\u4e00-\u9fff]{2,4})(表[。，])\s*(\d+)\.\s+([A-Z]{2}\d+)$', text)
        if m:
            action = m.group(1)
            punct_char = m.group(2)[1]
            num = m.group(3)
            code = m.group(4)
            text = f"{num}. {action}{code}表{punct_char}"
            b.text = text
            continue

        m = re.search(r'([。，])\s*(\d+)\.\s+([A-Z]{2}\d+)', text)
        if m:
            num = m.group(2)
            code = m.group(3)
            before = text[:m.start()].strip()
            after = text[m.end():].strip()
            if code in before:
                text = f"{num}. {before}"
            else:
                text = f"{num}. {before}{code}"
            if after and after.strip() != code:
                text += " " + after.strip()
            b.text = text
            continue

        m = re.search(r'([\u4e00-\u9fff])\s+(\d+)\.\s+([A-Z]{2}\d+)', text)
        if m:
            num = m.group(2)
            code = m.group(3)
            before = text[:m.start()].strip()
            last_char = m.group(1)
            after = text[m.end():].strip()
            if code in before:
                text = f"{num}. {before}{last_char}"
            else:
                text = f"{num}. {before}{last_char}{code}"
            if after and after.strip() != code:
                text += " " + after.strip()
            b.text = text
            continue

        m = re.search(r'([。，])\s*(\d+)\.\s*$', text)
        if m:
            num = m.group(2)
            before = text[:m.start()].strip()
            text = f"{num}. {before}"
            b.text = text
            continue

        m = re.search(r'([\u4e00-\u9fff])\s+(\d+)\.\s*$', text)
        if m:
            num = m.group(2)
            last_char = m.group(1)
            before = text[:m.start()].strip()
            text = f"{num}. {before}{last_char}"
            b.text = text

    return blocks


_HEADING_NUM_RE = re.compile(r'^(\d+(?:\.\d+)+)\s+(\S.*?)\s*$')
_HEADING_SWAP_RE = re.compile(r'^(.+?)[ 　]*(\d+\.\d+(?:\.\d+)*)$')
_HEADING_BAD_TAIL_RE = re.compile(r'[。？！?!.;；]\s*$')
_HEADING_BAD_INNER_RE = re.compile(r'。')
_HEADING_MAX_TITLE_LEN = 60


def _heading_level(num: str) -> int:
    return num.count('.') + 1


def _looks_like_heading_title(title: str) -> bool:
    if not title:
        return False
    if len(title) > _HEADING_MAX_TITLE_LEN:
        return False
    if _HEADING_BAD_TAIL_RE.search(title):
        return False
    if _HEADING_BAD_INNER_RE.search(title):
        return False
    return True


def fix_headings(blocks, options: Optional[ConvertOptions] = None):
    """修复章节标题格式并添加 Markdown 标题级别前缀"""
    opts = _opts(options)
    if not opts.fix_headings:
        return blocks

    for b in blocks:
        if b.type != "paragraph":
            continue
        text = b.text.strip()
        if not text:
            continue

        m = _HEADING_SWAP_RE.match(text)
        if m:
            title = m.group(1).strip()
            num = m.group(2).strip()
            text = f"{num} {title}"
            b.text = text

        m = _HEADING_NUM_RE.match(text)
        if m:
            num = m.group(1)
            title = m.group(2).strip()
            level = _heading_level(num)
            if 2 <= level <= 6 and _looks_like_heading_title(title):
                b.text = f"{'#' * level} {num} {title}"

    return blocks


def headers_match(a, b) -> bool:
    a_clean = [h.strip().lower() for h in a if h.strip()]
    b_clean = [h.strip().lower() for h in b if h.strip()]
    return a_clean == b_clean


def _header_similarity(a, b) -> float:
    def chars(headers):
        s = set()
        for h in headers:
            for c in h.strip().lower():
                if not c.isspace():
                    s.add(c)
        return s

    ca, cb = chars(a), chars(b)
    if not ca and not cb:
        return 1.0
    if not ca or not cb:
        return 0.0
    return len(ca & cb) / len(ca | cb)


_CODE_PATTERN_RE = re.compile(r'[A-Za-z]{2}\d{3,}')


def _looks_like_data_row(block: Block) -> bool:
    if not block.rows:
        return False
    first = block.rows[0]
    non_empty = [c.strip() for c in first if c.strip()]
    if not non_empty:
        return False
    for t in non_empty:
        if _CODE_PATTERN_RE.match(t):
            return True
    return False


def merge_consecutive_tables(blocks, options: Optional[ConvertOptions] = None):
    """合并连续出现的表格（跨页表格）"""
    opts = _opts(options)
    if not opts.merge_cross_page_tables:
        return blocks

    if len(blocks) < 2:
        return blocks

    result = []
    i = 0
    while i < len(blocks):
        current = blocks[i]
        if current.type != "table" or i + 1 >= len(blocks):
            result.append(current)
            i += 1
            continue

        nxt = blocks[i + 1]
        if nxt.type != "table":
            result.append(current)
            i += 1
            continue

        should_merge = False
        merged_rows = nxt.rows

        if headers_match(current.headers, nxt.headers):
            should_merge = True
        elif _looks_like_data_row(nxt):
            should_merge = True
        elif _header_similarity(current.headers, nxt.headers) >= 0.8:
            if _looks_like_data_row(nxt) or not headers_match(nxt.headers, nxt.rows[0] if nxt.rows else []):
                should_merge = True
            if should_merge and nxt.rows and not _looks_like_data_row(nxt):
                if _header_similarity(nxt.headers, list(nxt.rows[0])) >= 0.8:
                    merged_rows = nxt.rows[1:]

        if should_merge:
            merged = Block(
                type="table",
                y_top=current.y_top,
                headers=current.headers,
                rows=current.rows + merged_rows,
            )
            i += 2
            while i < len(blocks) and blocks[i].type == "table" and headers_match(current.headers, blocks[i].headers):
                merged.rows += blocks[i].rows
                i += 1
            result.append(merged)
        else:
            result.append(current)
            i += 1
    return result


# ══════════════════════════════════════════════════════════════════
# 表格数据处理（table_utils）
# ══════════════════════════════════════════════════════════════════

_CODE_PATTERN = re.compile(r'[A-Za-z]{2}\d{3,}')


def process_table_data(data, options: Optional[ConvertOptions] = None):
    """处理表格数据：扁平化多级表头、压缩空列、清理换行"""
    opts = _opts(options)
    if not data:
        return [], []

    def is_header_row(row):
        texts = [str(c or "").replace("\n", " ").strip() for c in row]
        non_empty = [t for t in texts if t]
        if not non_empty:
            return True
        for t in non_empty:
            if _CODE_PATTERN.match(t):
                return False
        if any(t.isdigit() for t in non_empty):
            return False
        if len(non_empty) >= 3:
            avg = sum(len(t) for t in non_empty) / len(non_empty)
            return avg <= 5
        return True

    data_start = 0
    for ri, row in enumerate(data):
        if is_header_row(row):
            data_start = ri + 1
        else:
            break

    header_rows = data[:data_start]
    data_rows = data[data_start:]

    if not data_rows:
        col_count = max(len(r) for r in header_rows) if header_rows else 0
        col_has = [False] * col_count
        for row in header_rows:
            for ci in range(min(len(row), col_count)):
                if row[ci] and str(row[ci]).strip():
                    col_has[ci] = True
        keep = [i for i, h in enumerate(col_has) if h]
        trimmed = []
        for row in header_rows:
            trimmed.append([str(row[i] or "") if i < len(row) else "" for i in keep])

        def clean(s: str) -> str:
            return s.replace("\n", "")

        flat = []
        for ci in range(len(trimmed[0])):
            parts = [clean(trimmed[ri][ci]).strip() for ri in range(len(trimmed))]
            flat.append(" ".join(p for p in parts if p))
        return flat, []

    col_count = max(
        max((len(r) for r in data_rows), default=0),
        max((len(r) for r in header_rows), default=0),
    )
    col_has = [False] * col_count
    for row in list(header_rows) + list(data_rows):
        for ci in range(min(len(row), col_count)):
            if row[ci] and str(row[ci]).strip():
                col_has[ci] = True
    keep = [i for i in range(col_count) if col_has[i]]
    if not keep:
        return [], []

    trimmed_rows = []
    for row in data_rows:
        trimmed_rows.append([str(row[i] or "") if i < len(row) else "" for i in keep])

    trimmed_headers = []
    for row in header_rows:
        trimmed_headers.append([str(row[i] or "") if i < len(row) else "" for i in keep])

    def clean_cell(s: str) -> str:
        return remove_cjk_spaces(s.replace("\n", " "))

    flat_headers = []
    for ci in range(len(keep)):
        parts = []
        for ri in range(len(trimmed_headers)):
            val = clean_cell(trimmed_headers[ri][ci]).strip()
            if val:
                parts.append(val)
        flat_headers.append(" ".join(parts))

    rows = [[clean_cell(c) for c in row] for row in trimmed_rows]

    flat_headers = _align_headers_to_data(flat_headers, rows)

    if opts.infer_empty_headers:
        for ci in range(len(flat_headers)):
            if flat_headers[ci].strip():
                continue
            data_vals = [r[ci] for r in rows if ci < len(r) and r[ci].strip()]
            if data_vals:
                unique_vals = set(v.strip() for v in data_vals)
                if len(unique_vals) == 1:
                    flat_headers[ci] = list(unique_vals)[0]

    final_headers = []
    final_rows = [[] for _ in range(len(rows))]
    for ci in range(len(flat_headers)):
        h = flat_headers[ci]
        if h.strip():
            final_headers.append(h)
            for ri in range(len(rows)):
                final_rows[ri].append(rows[ri][ci])
        else:
            has_data = any(rows[ri][ci].strip() for ri in range(len(rows)))
            if has_data:
                final_headers.append(h)
                for ri in range(len(rows)):
                    final_rows[ri].append(rows[ri][ci])

    return final_headers, final_rows


def _align_headers_to_data(headers, rows):
    """修复 PDF 合并单元格导致的表头/数据列错位"""
    if not headers or not rows:
        return headers
    n = len(headers)
    new_headers = list(headers)
    has_header = [bool(h.strip()) for h in headers]
    has_data = [
        any(ci < len(row) and str(row[ci]).strip() for row in rows)
        for ci in range(n)
    ]

    for ci in range(n):
        if has_header[ci] and not has_data[ci]:
            target = None
            if ci - 1 >= 0 and has_data[ci - 1] and not has_header[ci - 1]:
                target = ci - 1
            elif ci + 1 < n and has_data[ci + 1] and not has_header[ci + 1]:
                target = ci + 1
            if target is not None:
                moved = new_headers[ci].strip()
                existing = new_headers[target].strip()
                new_headers[target] = f"{moved} {existing}".strip() if existing else moved
                new_headers[ci] = ""
                has_header[ci] = False
                has_header[target] = True
    return new_headers


# ══════════════════════════════════════════════════════════════════
# Markdown 渲染（markdown_writer）
# ══════════════════════════════════════════════════════════════════

_BLANK_LINES_RE = re.compile(r"\n{3,}")


def render(blocks, title: str = "", options: Optional[ConvertOptions] = None,
           metadata: Optional[dict] = None) -> str:
    """将 Block 对象列表渲染为 Markdown 字符串"""
    opts = options if options is not None else ConvertOptions()
    lines = []

    if opts.add_front_matter and metadata:
        lines.append("---")
        for key, val in metadata.items():
            lines.append(f"{key}: {val}")
        lines.append("---")
        lines.append("")

    if title:
        lines.append(f"# {title}")
        lines.append("")

    prev_was_table = False
    for block in blocks:
        if block.type == "paragraph":
            text = block.text.strip()
            if not text:
                continue
            if prev_was_table:
                lines.append("")
            lines.append(text)
            prev_was_table = False
        elif block.type == "table":
            if lines and lines[-1] != "":
                lines.append("")
            lines.extend(_render_table(block))
            prev_was_table = True
        elif block.type == "image":
            if lines and lines[-1] != "":
                lines.append("")
            lines.append(f"![]({block.image_path})")
            lines.append("")
            prev_was_table = False

    content = "\n".join(lines).strip()

    if opts.merge_blank_lines:
        content = _BLANK_LINES_RE.sub("\n\n", content)

    return content


def _render_table(block: Block):
    """将 Table Block 渲染为 GFM 表格语法"""
    if not block.headers and not block.rows:
        return []

    def escape(s: str) -> str:
        return s.replace("|", "\\|").replace("\n", "")

    result = []

    header_line = "| " + " | ".join(escape(h) for h in block.headers) + " |"
    result.append(header_line)

    separator = "| " + " | ".join("---" for _ in block.headers) + " |"
    result.append(separator)

    for row in block.rows:
        padded = list(row) + [""] * (len(block.headers) - len(row))
        data_line = "| " + " | ".join(escape(padded[i]) for i in range(len(block.headers))) + " |"
        result.append(data_line)

    return result


# ══════════════════════════════════════════════════════════════════
# 自动编号解析（numbering）
#
# Word 的自动编号（多级列表 / 章节编号）并不写在正文文本里，而是由
# numbering.xml 定义、渲染时才显示。因此直接抽取 w:t 只能拿到 "概述"，
# 序号 "1" 会丢失。这里重建这套编号计算逻辑，把序号补回文本。
# ══════════════════════════════════════════════════════════════════

_ROMAN_UNITS = ((1000, 'm'), (900, 'cm'), (500, 'd'), (400, 'cd'), (100, 'c'),
                (90, 'xc'), (50, 'l'), (40, 'xl'), (10, 'x'), (9, 'ix'),
                (5, 'v'), (4, 'iv'), (1, 'i'))
_CN_DIGITS = '零一二三四五六七八九'
_CIRCLED_DIGITS = '①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳'
_LEGAL_DIGITS = '零壹贰叁肆伍陆柒捌玖'

_BULLET_FORMATS = ('bullet', 'none')


def _to_roman(n: int) -> str:
    if n <= 0:
        return str(n)
    out = []
    for value, sym in _ROMAN_UNITS:
        while n >= value:
            out.append(sym)
            n -= value
    return "".join(out)


def _to_letter(n: int) -> str:
    """1 -> A, 26 -> Z, 27 -> AA"""
    if n <= 0:
        return str(n)
    out = []
    while n > 0:
        n, rem = divmod(n - 1, 26)
        out.append(chr(ord('A') + rem))
    return "".join(reversed(out))


def _to_chinese(n: int) -> str:
    """阿拉伯数字转中文数字（chineseCounting 规则：10 写作「十」，11 写作「十一」）"""
    if n <= 0:
        return str(n)
    if n < 10:
        return _CN_DIGITS[n]
    if n < 20:
        return '十' + (_CN_DIGITS[n - 10] if n > 10 else '')
    if n < 100:
        rem = n % 10
        return _CN_DIGITS[n // 10] + '十' + (_CN_DIGITS[rem] if rem else '')
    if n < 1000:
        return _CN_DIGITS[n // 100] + '百' + _cn_tail(n % 100, _to_chinese, '一')
    if n < 10000:
        return _CN_DIGITS[n // 1000] + '千' + _cn_tail(n % 1000, _to_chinese, '一', True)
    return str(n)


def _cn_tail(rem: int, render, lead: str = '', zero_lead: bool = False) -> str:
    """千位/百位之后的部分：被跳过的位补「零」，小写 10~19 补前导「一」（大写「壹拾」自带）"""
    if rem == 0:
        return ''
    if rem < 10:
        return '零' + render(rem)
    zero = '零' if zero_lead and rem < 100 else ''
    text = render(rem)
    if lead and 10 <= rem < 20:
        text = lead + text
    return zero + text


def _to_chinese_legal(n: int) -> str:
    """大写中文数字（chineseLegalSimplified）"""
    if n <= 0:
        return str(n)
    if n < 10:
        return _LEGAL_DIGITS[n]
    if n < 100:
        rem = n % 10
        return _LEGAL_DIGITS[n // 10] + '拾' + (_LEGAL_DIGITS[rem] if rem else '')
    if n < 1000:
        return _LEGAL_DIGITS[n // 100] + '佰' + _cn_tail(n % 100, _to_chinese_legal)
    if n < 10000:
        return _LEGAL_DIGITS[n // 1000] + '仟' + _cn_tail(n % 1000, _to_chinese_legal, '', True)
    return str(n)


def _format_num_value(value: int, fmt: str) -> str:
    """按 w:numFmt 渲染单个级别的编号值"""
    fmt = (fmt or 'decimal').strip()
    if fmt == 'decimal':
        return str(value)
    if fmt == 'decimalZero':
        return str(value).zfill(2)
    if fmt == 'upperRoman':
        return _to_roman(value).upper()
    if fmt == 'lowerRoman':
        return _to_roman(value).lower()
    if fmt == 'upperLetter':
        return _to_letter(value).upper()
    if fmt == 'lowerLetter':
        return _to_letter(value).lower()
    if fmt == 'decimalFullWidth':
        return "".join(chr(ord(c) + 0xFEE0) for c in str(value))
    if fmt == 'decimalHalfWidth':
        return str(value)
    if fmt == 'chineseCounting' or fmt == 'chineseCountingThousand':
        return _to_chinese(value)
    if fmt == 'chineseLegalSimplified':
        return _to_chinese_legal(value)
    if fmt == 'ideographDigital':
        return _to_chinese(value)
    if fmt == 'decimalEnclosedCircle':
        return _CIRCLED_DIGITS[value - 1] if 1 <= value <= len(_CIRCLED_DIGITS) else f"({value})"
    if fmt == 'decimalEnclosedParen':
        return f"({value})"
    if fmt == 'decimalEnclosedFullstop':
        return f"{value}."
    if fmt == 'hex':
        return f"{value:X}"
    return str(value)


class NumberingResolver:
    """按文档顺序重建 Word 自动编号（numPr + numbering.xml）"""

    def __init__(self, doc=None):
        self._doc = doc
        self._abstracts = {}       # abstractNumId -> {ilvl: lvl定义}
        self._abstract_links = {}  # abstractNumId -> {styleLink, numStyleLink}
        self._num_map = {}         # numId -> abstractNumId
        self._start_override = {}  # (numId, ilvl) -> 起始值
        self._lvl_override = {}    # (numId, ilvl) -> lvl定义（lvlOverride 内联定义）
        self._pstyle_map = {}      # 样式名/styleId -> (numId, ilvl)
        self._counters = {}        # numId -> {ilvl: 当前值}
        self._depth = {}           # numId -> 当前展开的层级深度
        self._last_level = {}      # numId -> 上一段落的 ilvl
        self._style_cache = {}     # styleId -> (numId, ilvl)
        self._levels_cache = {}    # numId -> 级别定义
        self._loaded = False

    # ── 加载 ────────────────────────────────────────────────────────
    def load(self) -> "NumberingResolver":
        if self._loaded:
            return self
        self._loaded = True
        if self._doc is None:
            return self
        numbering = None
        try:
            part = self._doc.part.numbering_part
            numbering = part.element if part is not None else None
        except Exception:
            numbering = None
        if numbering is None:
            return self

        for child in numbering:
            if child.tag == qn('w:abstractNum'):
                self._parse_abstract(child)
            elif child.tag == qn('w:num'):
                self._parse_num(child)

        # 「将级别链接到样式」：编号级别上的 w:pStyle
        for num_id in sorted(self._num_map, key=lambda v: (_to_int(v) is None, _to_int(v) or 0)):
            levels = self._levels(num_id)
            for ilvl in sorted(k for k in levels if isinstance(k, int)):
                style_ref = (levels[ilvl] or {}).get('pStyle')
                if not style_ref:
                    continue
                for key in self._pstyle_keys(style_ref):
                    self._pstyle_map.setdefault(key, (num_id, ilvl))
        return self

    def _pstyle_keys(self, style_ref):
        """w:pStyle 记录的是 styleId，这里同时收集样式名，便于两种方式匹配"""
        ref = (style_ref or '').strip()
        keys = {ref.lower()}
        st = self._style_elem(ref)
        if st is None:
            st = self._style_elem_by_name(ref)
        if st is not None:
            sid = st.get(qn('w:styleId'))
            if sid:
                keys.add(sid.lower())
            nm = st.find(qn('w:name'))
            if nm is not None and nm.get(qn('w:val')):
                keys.add(nm.get(qn('w:val')).strip().lower())
            try:
                keys.add(BabelFish.internal2ui(ref).strip().lower())
            except Exception:
                pass
        return {k for k in keys if k}

    def _parse_abstract(self, elem):
        aid = elem.get(qn('w:abstractNumId'))
        if aid is None:
            return
        levels = {}
        for lvl in elem.findall(qn('w:lvl')):
            ilvl = _to_int(lvl.get(qn('w:ilvl')))
            if ilvl is None:
                continue
            levels[ilvl] = _parse_lvl(lvl)
        self._abstracts[aid] = levels

        style_link = elem.find(qn('w:styleLink'))
        num_style_link = elem.find(qn('w:numStyleLink'))
        self._abstract_links[aid] = {
            'styleLink': style_link.get(qn('w:val')) if style_link is not None else None,
            'numStyleLink': (num_style_link.get(qn('w:val'))
                             if num_style_link is not None else None),
        }

    def _parse_num(self, elem):
        nid = elem.get(qn('w:numId'))
        aid = elem.find(qn('w:abstractNumId'))
        if nid is None or aid is None:
            return
        self._num_map[nid] = aid.get(qn('w:val'))
        for ov in elem.findall(qn('w:lvlOverride')):
            ilvl = _to_int(ov.get(qn('w:ilvl')))
            if ilvl is None:
                continue
            so = ov.find(qn('w:startOverride'))
            if so is not None:
                start = _to_int(so.get(qn('w:val')))
                if start is not None:
                    self._start_override[(nid, ilvl)] = start
            inline_lvl = ov.find(qn('w:lvl'))
            if inline_lvl is not None:
                self._lvl_override[(nid, ilvl)] = _parse_lvl(inline_lvl)

    # ── 样式 → 编号 ─────────────────────────────────────────────────
    def _styles_root(self):
        try:
            return self._doc.styles.element
        except Exception:
            return None

    def _style_elem_by_name(self, name):
        root = self._styles_root()
        if root is None or not name:
            return None
        target = (name or '').strip().lower()
        for st in root.findall(qn('w:style')):
            nm = st.find(qn('w:name'))
            if nm is None:
                continue
            if (nm.get(qn('w:val')) or '').strip().lower() == target:
                return st
        return None

    def _style_elem(self, style_id):
        root = self._styles_root()
        if root is None or not style_id:
            return None
        for st in root.findall(qn('w:style')):
            if st.get(qn('w:styleId')) == style_id:
                return st
        return None

    def _style_numpr(self, style_id, depth=0):
        """沿 basedOn 链查找样式自身的编号设置"""
        if style_id is None or depth > 10:
            return None, None
        if style_id in self._style_cache:
            return self._style_cache[style_id]
        self._style_cache[style_id] = (None, None)  # 防止循环引用
        st = self._style_elem(style_id)
        num_id = ilvl = None
        if st is not None:
            ppr = st.find(qn('w:pPr'))
            numpr = ppr.find(qn('w:numPr')) if ppr is not None else None
            if numpr is not None:
                e = numpr.find(qn('w:numId'))
                if e is not None:
                    num_id = e.get(qn('w:val'))
                e = numpr.find(qn('w:ilvl'))
                if e is not None:
                    ilvl = e.get(qn('w:val'))
            if num_id is None:
                based = st.find(qn('w:basedOn'))
                if based is not None:
                    num_id, ilvl = self._style_numpr(based.get(qn('w:val')), depth + 1)
        result = (num_id, ilvl)
        self._style_cache[style_id] = result
        return result

    def _pstyle_num(self, style_id, style_name):
        """编号定义中通过 w:pStyle 关联到样式的级别（Word 的「将级别链接到样式」）"""
        for key in (style_id, style_name):
            if not key:
                continue
            found = self._pstyle_map.get(key.strip().lower())
            if found is not None:
                return found
        return None

    # ── 段落 → (numId, ilvl) ────────────────────────────────────────
    def _paragraph_numpr(self, para):
        num_id = ilvl = None
        ppr = para._element.find(qn('w:pPr'))
        numpr = ppr.find(qn('w:numPr')) if ppr is not None else None
        if numpr is not None:
            e = numpr.find(qn('w:numId'))
            if e is not None:
                num_id = e.get(qn('w:val'))
            e = numpr.find(qn('w:ilvl'))
            if e is not None:
                ilvl = e.get(qn('w:val'))

        style_name = None
        style_id = None
        try:
            style = para.style
        except Exception:
            style = None
        if style is not None:
            style_name = getattr(style, "name", None)
            style_id = getattr(style, "style_id", None)
            if style_id is None:
                el = getattr(style, "_element", None)
                style_id = el.get(qn('w:styleId')) if el is not None else None

        if num_id is None and style_id is not None:
            s_num, s_ilvl = self._style_numpr(style_id)
            num_id, ilvl = s_num, s_ilvl

        if num_id == '0':
            # 显式声明 numId=0 表示取消编号，不再回退到样式/pStyle 关联
            return None, None

        if not num_id:
            # 段落与样式都没直接声明时，尝试按 numbering 的 pStyle 关联反查
            linked = self._pstyle_num(style_id, style_name)
            if linked is None:
                return None, None
            num_id, ilvl = linked[0], linked[1]

        return num_id, (_to_int(ilvl) or 0 if ilvl is not None else 0)

    # ── 编号计算 ────────────────────────────────────────────────────
    def _levels(self, num_id, _seen=None):
        if num_id in self._levels_cache:
            return self._levels_cache[num_id]
        _seen = _seen if _seen is not None else set()
        aid = self._num_map.get(num_id)
        levels = self._abstracts.get(aid) if aid is not None else None
        links = self._abstract_links.get(aid, {}) if aid is not None else {}
        link_name = links.get('numStyleLink')
        if levels and link_name and link_name not in _seen:
            # 编号样式（w:numStyleLink）指向真正带 styleLink 的编号定义
            _seen.add(link_name)
            st = self._style_elem_by_name(link_name)
            if st is not None:
                link_num, _ = self._style_numpr(st.get(qn('w:styleId')))
                if link_num:
                    linked = self._levels(link_num, _seen)
                    if linked:
                        levels = linked
        self._levels_cache[num_id] = levels
        return levels

    def _lvl_def(self, num_id, levels, ilvl):
        override = self._lvl_override.get((num_id, ilvl))
        base = levels.get(ilvl) if levels else None
        if override and base:
            merged = dict(base)
            merged.update(override)
            return merged
        return override or base

    def _start_value(self, num_id, levels, ilvl):
        override = self._start_override.get((num_id, ilvl))
        if override is not None:
            return override
        lvl = self._lvl_def(num_id, levels, ilvl)
        if lvl:
            return lvl.get('start', 1) or 1
        return 1

    def _next_value(self, num_id, levels, ilvl):
        """推进计数器并返回该级别的当前值（同时保持各级别上下文）"""
        counters = self._counters.setdefault(num_id, {})
        depth = self._depth.get(num_id, 0)
        start = self._start_value(num_id, levels, ilvl)

        if ilvl > depth:
            for lv in range(depth, ilvl):
                if counters.get(lv) is None:
                    counters[lv] = self._start_value(num_id, levels, lv)
            depth = ilvl
        elif ilvl < depth:
            for lv in range(ilvl + 1, depth):
                counters.pop(lv, None)
            depth = ilvl

        lvl = self._lvl_def(num_id, levels, ilvl)
        restart = (lvl or {}).get('restart')
        last = self._last_level.get(num_id)
        if restart is not None and restart < ilvl and last is not None and last <= restart:
            counters.pop(ilvl, None)

        current = counters.get(ilvl)
        counters[ilvl] = start if current is None else current + 1
        self._depth[num_id] = ilvl + 1
        self._last_level[num_id] = ilvl
        return counters[ilvl]

    def _render_lvl_text(self, text, num_id, levels, ilvl):
        counters = self._counters.get(num_id, {})

        def repl(m):
            if m.group(0) == '%%':  # 转义后的字面量 %
                return '%'
            idx = int(m.group(1)) - 1
            if idx < 0:
                return m.group(0)
            value = counters.get(idx)
            if value is None:
                value = self._start_value(num_id, levels, idx)
            fmt = 'decimal'
            lvl = self._lvl_def(num_id, levels, idx)
            if lvl:
                fmt = lvl.get('fmt', 'decimal')
            return _format_num_value(value, fmt)

        return _LVL_PLACEHOLDER_RE.sub(repl, text)

    # ── 对外接口 ────────────────────────────────────────────────────
    def number_for(self, para) -> str:
        """返回段落自动编号文本（无编号时返回空串，并已推进内部计数器）"""
        return self._resolve(para)[0]

    def prefix_for(self, para) -> str:
        """返回可直接拼在段落文本前的编号前缀（含分隔符），无编号时为空串"""
        number, suff = self._resolve(para)
        if not number:
            return ""
        if suff == 'nothing' or number.endswith((' ', '\u3000')):
            return number
        # 中文标点后不再补空格，保持与 Word 中「1、标题」的排版一致
        if number[-1] in '、，。；：？！）】》':
            return number
        return f"{number} "

    def _resolve(self, para):
        """计算段落编号，返回 (编号文本, 分隔方式)"""
        num_id, ilvl = self._paragraph_numpr(para)
        if num_id is None:
            return "", 'tab'
        levels = self._levels(num_id)
        if not levels:
            return "", 'tab'
        lvl = self._lvl_def(num_id, levels, ilvl)
        if not lvl:
            return "", 'tab'
        if (lvl.get('fmt') or 'decimal') in _BULLET_FORMATS:
            # 项目符号类编号不还原为文本
            return "", lvl.get('suff', 'tab')
        self._next_value(num_id, levels, ilvl)
        text = lvl.get('text') or ''
        if not text:
            return "", lvl.get('suff', 'tab')
        return self._render_lvl_text(text, num_id, levels, ilvl), lvl.get('suff', 'tab')


_LVL_PLACEHOLDER_RE = re.compile(r'%%|%(\d)')


def _to_int(value):
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return None


def _parse_lvl(lvl):
    """解析单个 w:lvl 定义"""

    def attr(tag, default=None):
        e = lvl.find(qn(tag))
        if e is None:
            return default
        val = e.get(qn('w:val'))
        return default if val is None else val

    return {
        'start': _to_int(attr('w:start', '1')) or 1,
        'fmt': attr('w:numFmt', 'decimal') or 'decimal',
        'text': attr('w:lvlText', '') or '',
        'restart': _to_int(attr('w:lvlRestart')),
        'pStyle': attr('w:pStyle'),
        'suff': attr('w:suff', 'tab') or 'tab',
    }


# ══════════════════════════════════════════════════════════════════
# Word 解析（word_engine）
# ══════════════════════════════════════════════════════════════════

HEADING_LEVEL_MAP = {
    "Heading 1": 1, "Heading 2": 2, "Heading 3": 3,
    "Heading 4": 4, "Heading 5": 5, "Heading 6": 6,
    "标题 1": 1, "标题 2": 2, "标题 3": 3,
    "标题 4": 4, "标题 5": 5, "标题 6": 6,
}


def parse_docx(path: str, options: Optional[ConvertOptions] = None) -> list:
    """解析 .docx 文件，返回按文档顺序排列的 Block 列表"""
    doc = Document(path)
    blocks = []
    order = 0

    numbering = NumberingResolver(doc).load()

    body = doc.element.body
    para_idx = 0
    tbl_idx = 0
    list_counter = 0

    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            if para_idx >= len(doc.paragraphs):
                break
            para = doc.paragraphs[para_idx]
            para_idx += 1
            text = _render_paragraph(para, list_counter, doc, numbering)
            if text.strip():
                style_name = para.style.name if para.style else ""
                if HEADING_LEVEL_MAP.get(style_name):
                    # 标题不参与正文列表计数
                    list_counter = 0
                elif _is_numbered_list(para):
                    list_counter += 1
                elif _is_bullet_list(para):
                    pass
                else:
                    list_counter = 0
                blocks.append(Block(type="paragraph", y_top=order, text=text))
                order += 1
        elif child.tag == qn('w:tbl'):
            if tbl_idx >= len(doc.tables):
                break
            tbl = doc.tables[tbl_idx]
            tbl_idx += 1
            data = _extract_table_grid(tbl)
            headers, rows = process_table_data(data, options)
            if headers and rows:
                blocks.append(Block(
                    type="table", y_top=order,
                    headers=headers, rows=rows,
                ))
                order += 1

    blocks = merge_consecutive_tables(blocks, options)
    blocks = merge_broken_paragraphs(blocks, options)
    blocks = fix_list_number_positions(blocks, options)
    blocks = fix_headings(blocks, options)
    return blocks


def _render_paragraph(para, list_counter: int, doc=None, numbering=None) -> str:
    """渲染段落为 Markdown 文本（含样式前缀、Word 自动编号与行内格式）"""
    style_name = para.style.name if para.style else ""

    inline = _render_inline(para, doc)

    level = HEADING_LEVEL_MAP.get(style_name)
    if level:
        # Word 的自动编号不在文本里，需按 numbering.xml 还原后补在标题前
        prefix = numbering.prefix_for(para) if numbering else ""
        if prefix:
            inline = prefix + inline.lstrip()
        return f"{'#' * level} {inline}"

    if _is_bullet_list(para):
        return f"- {inline}"
    if _is_numbered_list(para):
        prefix = numbering.prefix_for(para) if numbering else ""
        if prefix:
            return prefix + inline.lstrip()
        return f"{list_counter + 1}. {inline}"

    return inline


def _render_inline(para, doc=None) -> str:
    """渲染段落的行内内容：run 的加粗/斜体/删除线 + 超链接"""
    parts = []
    for child in para._element:
        if child.tag == qn('w:r'):
            parts.append(_render_run(child))
        elif child.tag == qn('w:hyperlink'):
            parts.append(_render_hyperlink(child, doc))
    return "".join(parts)


def _render_run(r_element) -> str:
    """渲染单个 run 元素，应用加粗/斜体/删除线"""
    texts = []
    for t in r_element.iter():
        if t.tag == qn('w:t'):
            texts.append(t.text or "")
        elif t.tag == qn('w:br') or t.tag == qn('w:cr'):
            texts.append("\n")
        elif t.tag == qn('w:tab'):
            texts.append("\t")
    text = "".join(texts)
    if not text:
        return ""

    rpr = r_element.find(qn('w:rPr'))
    bold = italic = strike = False
    if rpr is not None:
        if rpr.find(qn('w:b')) is not None:
            bold = True
        if rpr.find(qn('w:i')) is not None:
            italic = True
        if rpr.find(qn('w:strike')) is not None:
            strike = True

    if strike:
        text = f"~~{text}~~"
    if bold:
        text = f"**{text}**"
    if italic:
        text = f"*{text}*"
    return text


def _render_hyperlink(hl_element, doc=None) -> str:
    """渲染超链接元素为 [text](url)"""
    texts = []
    for r in hl_element.iter(qn('w:r')):
        texts.append(_render_run(r))
    text = "".join(texts)

    rid = hl_element.get(qn('r:id'))
    url = ""
    if rid and doc is not None:
        try:
            rel = doc.part.rels.get(rid)
            if rel is not None:
                url = rel.target_ref or ""
        except Exception:
            url = ""
    if url:
        return f"[{text}]({url})"
    return text


def _is_bullet_list(para) -> bool:
    """判断是否为无序列表项"""
    style_name = para.style.name if para.style else ""
    if "List Bullet" in style_name or "项目符号" in style_name:
        return True
    ppr = para._element.find(qn('w:pPr'))
    if ppr is not None:
        numpr = ppr.find(qn('w:numPr'))
        if numpr is not None:
            numfmt = numpr.find(qn('w:numFmt'))
            if numfmt is not None and numfmt.get(qn('w:val')) == 'bullet':
                return True
    return False


def _is_numbered_list(para) -> bool:
    """判断是否为有序列表项"""
    style_name = para.style.name if para.style else ""
    if "List Number" in style_name or "编号" in style_name:
        return True
    ppr = para._element.find(qn('w:pPr'))
    if ppr is not None:
        numpr = ppr.find(qn('w:numPr'))
        if numpr is not None:
            numfmt = numpr.find(qn('w:numFmt'))
            if numfmt is not None and numfmt.get(qn('w:val')) in ('decimal', 'chineseCounting'):
                return True
    return False


def _extract_table_grid(tbl) -> list:
    """提取表格为二维网格，正确处理合并单元格"""
    grid = []
    vmerge_pending = {}

    for tr in tbl._element.iter(qn('w:tr')):
        row = []
        col_idx = 0
        tc_list = list(tr.findall(qn('w:tc')))
        tc_iter = iter(tc_list)

        while col_idx < _table_grid_cols(tbl) or tc_iter:
            try:
                tc = next(tc_iter)
            except StopIteration:
                break

            tcpr = tc.find(qn('w:tcPr'))
            grid_span = 1
            vmerge_state = "none"
            if tcpr is not None:
                gs = tcpr.find(qn('w:gridSpan'))
                if gs is not None:
                    try:
                        grid_span = int(gs.get(qn('w:val'), '1'))
                    except ValueError:
                        grid_span = 1
                vm = tcpr.find(qn('w:vMerge'))
                if vm is not None:
                    val = vm.get(qn('w:val'))
                    vmerge_state = "restart" if val == "restart" else "continue"

            cell_text = _tc_text(tc)

            if vmerge_state == "restart":
                for i in range(grid_span):
                    vmerge_pending[col_idx + i] = cell_text if i == 0 else ""
                row.append(cell_text)
                for _ in range(grid_span - 1):
                    row.append("")
                col_idx += grid_span
            elif vmerge_state == "continue":
                for i in range(grid_span):
                    row.append(vmerge_pending.get(col_idx + i, ""))
                col_idx += grid_span
            else:
                for i in range(grid_span):
                    vmerge_pending.pop(col_idx + i, None)
                row.append(cell_text)
                for _ in range(grid_span - 1):
                    row.append("")
                col_idx += grid_span

        if row:
            grid.append(row)

    if not grid:
        return [[cell.text for cell in row.cells] for row in tbl.rows]
    return grid


def _table_grid_cols(tbl) -> int:
    """获取表格的网格列数（tblGrid 的 gridCol 数量）"""
    tbl_grid = tbl._element.find(qn('w:tblGrid'))
    if tbl_grid is not None:
        return len(list(tbl_grid.iter(qn('w:gridCol'))))
    first_row = tbl._element.find(qn('w:tr'))
    if first_row is not None:
        return len(list(first_row.iter(qn('w:tc'))))
    return 0


def _tc_text(tc_element) -> str:
    """提取 tc 元素的纯文本"""
    texts = []
    for t in tc_element.iter(qn('w:t')):
        texts.append(t.text or "")
    return "".join(texts)


# ══════════════════════════════════════════════════════════════════
# 主入口
# ══════════════════════════════════════════════════════════════════

def convert(docx_path: str, options: Optional[ConvertOptions] = None) -> str:
    """将 docx 转换为 markdown 文本"""
    blocks = parse_docx(docx_path, options)
    return render(blocks, options=options)


def _convert_one(src: Path, opts: ConvertOptions) -> bool:
    """将单个 docx 文件转换为同目录同名的 md 文件，成功返回 True"""
    try:
        content = convert(str(src), opts)
    except Exception as e:
        print(f"转换失败: {src}: {e}", file=sys.stderr)
        return False

    out_path = src.with_suffix(".md")
    try:
        out_path.write_text(content, encoding="utf-8")
    except OSError as e:
        print(f"写入失败: {out_path}: {e}", file=sys.stderr)
        return False

    print(f"已生成: {out_path}")
    return True


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        prog="docx2md",
        description="将 Word (.docx) 文件转换为 Markdown，输出同目录同名 .md 文件",
    )
    parser.add_argument("path", nargs="?", help="docx 文件路径（与 -d 二选一）")
    parser.add_argument("-d", "--dir", metavar="DIR", help="将目录下所有 docx 文件转换为 md 文件")
    parser.add_argument("-b", "--business", action="store_true",
                        help="启用业务文档规则（序号修复、空表头推断）")
    args = parser.parse_args(argv)

    opts = ConvertOptions.business() if args.business else ConvertOptions.default()

    # 目录模式：转换目录下所有 docx
    if args.dir:
        dir_path = Path(args.dir)
        if not dir_path.is_dir():
            print(f"错误: 目录不存在: {dir_path}", file=sys.stderr)
            return 1
        docx_files = sorted(dir_path.glob("*.docx"))
        if not docx_files:
            print(f"目录中没有 .docx 文件: {dir_path}", file=sys.stderr)
            return 1
        success = 0
        for f in docx_files:
            if _convert_one(f, opts):
                success += 1
        print(f"完成: 成功 {success} / 共 {len(docx_files)} 个文件")
        return 0 if success == len(docx_files) else 1

    # 单文件模式
    if not args.path:
        parser.error("必须提供文件路径或 -d/--dir 目录")
    src = Path(args.path)
    if not src.exists():
        print(f"错误: 文件不存在: {src}", file=sys.stderr)
        return 1
    if src.suffix.lower() != ".docx":
        print(f"错误: 仅支持 .docx 文件: {src}", file=sys.stderr)
        return 1

    return 0 if _convert_one(src, opts) else 1


if __name__ == "__main__":
    sys.exit(main())
